import io
import math
import openpyxl
import pandas as pd
import streamlit as st
import zipfile
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

# ==============================================================================
# 1. ตั้งค่าและตกแต่งด้วย CSS
# ==============================================================================
st.set_page_config(
    page_title="Optics Lens Dispatcher System",
    page_icon="👓",
    layout="wide",
)

custom_css = """
<style>
    @import url('https://fonts.googleapis.com/css2?family=Kanit:wght@300;400;500;600&display=swap');
    
    html, body, [class*="css"] {
        font-family: 'Kanit', sans-serif;
    }

    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}

    .stApp {
        background: linear-gradient(180deg, #EBF1F7 0%, #E2E9F1 100%);
    }

    .header-box {
        background: #FFFFFF;
        padding: 25px 20px;
        border-radius: 20px;
        color: #1E293B;
        text-align: center;
        margin-bottom: 20px;
        box-shadow: 0 10px 25px rgba(148, 163, 184, 0.35),
                    inset 0 2px 3px rgba(255, 255, 255, 0.8);
        border: 1px solid rgba(255, 255, 255, 0.6);
    }

    .header-title {
        font-size: 26px;
        font-weight: 600;
        color: #1E3A8A;
        margin-bottom: 6px;
    }

    .header-subtitle {
        font-size: 14px;
        color: #64748B;
    }

    [data-testid="stFileUploader"] {
        background: #FFFFFF;
        border-radius: 16px;
        padding: 15px;
        box-shadow: 0 8px 20px rgba(148, 163, 184, 0.25);
        border: 2px dashed #93C5FD;
    }

    div.stButton > button:first-child {
        background: linear-gradient(135deg, #2563EB 0%, #1D4ED8 100%);
        color: white;
        font-size: 16px;
        font-weight: 500;
        border: none;
        border-radius: 14px;
        padding: 12px 24px;
        width: 100%;
        transition: all 0.2s ease-in-out;
        box-shadow: 0 6px 18px rgba(37, 99, 235, 0.35);
    }
    
    div.stButton > button:first-child:hover {
        transform: translateY(-2px);
        box-shadow: 0 10px 22px rgba(37, 99, 235, 0.45);
        background: linear-gradient(135deg, #1D4ED8 0%, #1E40AF 100%);
    }

    div.stDownloadButton > button:first-child {
        background: linear-gradient(135deg, #10B981 0%, #059669 100%);
        color: white;
        font-size: 16px;
        font-weight: 500;
        border: none;
        border-radius: 14px;
        padding: 12px 24px;
        width: 100%;
        transition: all 0.2s ease-in-out;
        box-shadow: 0 6px 18px rgba(16, 185, 129, 0.35);
    }
    
    div.stDownloadButton > button:first-child:hover {
        transform: translateY(-2px);
        box-shadow: 0 10px 22px rgba(16, 185, 129, 0.45);
        background: linear-gradient(135deg, #059669 0%, #047857 100%);
    }
</style>
"""
st.markdown(custom_css, unsafe_allow_html=True)

# ==============================================================================
# 2. คลังข้อความสองภาษา
# ==============================================================================
TEXTS = {
    "TH": {
        "title": "👓 Optics Lens Dispatcher System",
        "subtitle": "ระบบจัดกลุ่ม จัดเรียง และเตรียมหน้าพิมพ์ใบจัดส่งเลนส์อัตโนมัติ",
        "uploader_label": "ลากไฟล์ Excel ต้นฉบับ (.xlsx) มาวางที่นี่ (รองรับหลายไฟล์พร้อมกัน)",
        "btn_process": "🚀 ประมวลผลและแปลงไฟล์ทั้งหมด",
        "guide_title": "📖 คำแนะนำและวิธีการใช้งานระบบ",
        "guide_content": """
            **🔹 ขั้นตอนการทำงาน:**
            1. **อัปโหลดไฟล์:** ลากไฟล์รายงาน Excel ต้นฉบับ (.xlsx) มาวางในช่องด้านบน (สามารถเลือกพร้อมกันได้หลายไฟล์)
            2. **ประมวลผล:** กดปุ่ม **"🚀 ประมวลผลและแปลงไฟล์ทั้งหมด"** ระบบจะจัดกลุ่มสาขา เรียงลำดับ DO และตั้งค่าจัดหน้ากระดาษพิมพ์ให้อัตโนมัติ
            3. **ตรวจสอบ & ดาวน์โหลด:** ดูสถิติภาพรวมบนตาราง และดาวน์โหลดไฟล์ Excel สรุปผลไปสั่งพิมพ์ได้ทันที
            
            <hr style="margin: 10px 0;">
            
            **💡 สัญลักษณ์สีในไฟล์ Excel:**
            * 🟢 **Tab สีเขียว (Normal):** สั่งกด Print รวมทั้งไฟล์ได้ทันที (ไม่เกิน 8 DO และไม่เกิน 35 แถว)
            * 🔴 **Tab สีแดง (Heavy):** ต้องเปิดเข้าไปสั่ง Print แยกทีละ Sheet (เกิน 8 DO หรือเกิน 35 แถว)
        """,
        "success": "✅ **ประมวลผลสำเร็จเรียบร้อย!**",
        "dashboard_title": "📊 ภาพรวมการจัดส่ง (Interactive Dashboard)",
        "m_total_stores": "สาขาทั้งหมด",
        "m_heavy_stores": "🔴 พิมพ์แยก (Heavy)",
        "m_normal_stores": "🟢 พิมพ์รวม (Normal)",
        "m_total_pcs": "ยอดเลนส์รวม",
        "search_label": "🔍 ค้นหาสาขา (รหัสสาขา หรือ ชื่อสาขา):",
        "filter_label": "📌 กรองตามสถานะ:",
        "filter_all": "ทั้งหมด (All)",
        "filter_normal": "🟢 พิมพ์รวม (Normal)",
        "filter_heavy": "🔴 พิมพ์แยก (Heavy)",
        "main_export": "📦 ดาวน์โหลดชุดไฟล์หลัก",
        "single_export": "🚨 ดาวน์โหลดฉุกเฉินเฉพาะสาขา",
        "btn_download_excel": "📥 ดาวน์โหลดไฟล์ Excel สรุปผล",
        "btn_download_zip": "📦 ดาวน์โหลดไฟล์ทั้งหมดเป็น ZIP",
        "select_single": "เลือกสาขาที่ต้องการดาวน์โหลดไฟล์เดี่ยว:",
        "btn_single": "📄 ดาวน์โหลด Excel เฉพาะสาขา",
        "btn_reset": "🔄 อัปโหลดไฟล์ชุดใหม่",
        "unit_store": "สาขา",
        "unit_pcs": "ชิ้น",
        "status_normal": "🟢 พิมพ์รวม (Normal)",
        "status_heavy": "🔴 พิมพ์แยก (Heavy)",
    },
    "EN": {
        "title": "👓 Optics Lens Dispatcher System",
        "subtitle": "Automated Lens Dispatching, Sorting, and Print-Ready Processing System",
        "uploader_label": "Drag and drop original Excel files (.xlsx) here (Supports multiple files)",
        "btn_process": "🚀 Process & Convert All Files",
        "guide_title": "📖 System User Guide",
        "guide_content": """
            **🔹 Workflow Steps:**
            1. **Upload Files:** Drag & drop original Excel reports (.xlsx) into the drop zone above (batch upload supported).
            2. **Process Data:** Click **"🚀 Process & Convert All Files"** to auto-group stores, sort DOs, and format print pages.
            3. **Review & Download:** Inspect summary statistics and download print-ready Excel files immediately.
            
            <hr style="margin: 10px 0;">
            
            **💡 Sheet Color Indicators in Excel:**
            * 🟢 **Green Tabs (Normal):** Print entire workbook at once (≤ 8 DOs AND ≤ 35 rows).
            * 🔴 **Red Tabs (Heavy):** Must print each sheet individually (> 8 DOs OR > 35 rows).
        """,
        "success": "✅ **Processing Completed Successfully!**",
        "dashboard_title": "📊 Dispatch Summary (Interactive Dashboard)",
        "m_total_stores": "Total Stores",
        "m_heavy_stores": "🔴 Separate Print (Heavy)",
        "m_normal_stores": "🟢 Batch Print (Normal)",
        "m_total_pcs": "Total Lenses",
        "search_label": "🔍 Search Store (Store ID or Store Name):",
        "filter_label": "📌 Filter by Status:",
        "filter_all": "All Stores",
        "filter_normal": "🟢 Batch Print (Normal)",
        "filter_heavy": "🔴 Separate Print (Heavy)",
        "main_export": "📦 Download Main Package",
        "single_export": "🚨 Quick Single Store Download",
        "btn_download_excel": "📥 Download Summary Excel File",
        "btn_download_zip": "📦 Download All Files as ZIP",
        "select_single": "Select specific store to download:",
        "btn_single": "📄 Download Store File",
        "btn_reset": "🔄 Upload New Files",
        "unit_store": "Stores",
        "unit_pcs": "Pcs",
        "status_normal": "🟢 Batch Print (Normal)",
        "status_heavy": "🔴 Separate Print (Heavy)",
    },
}

# ==============================================================================
# 3. ฟังก์ชันประมวลผล Excel
# ==============================================================================
def parse_num(val):
    if pd.isna(val) or val == "" or val is None:
        return 0.0
    try:
        return float(val)
    except ValueError:
        return 0.0


def setup_sheet_page_layout(ws, is_heavy=False):
    ws.views.sheetView[0].showGridLines = True
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.blackAndWhite = True

    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.header = 0.1
    ws.page_margins.footer = 0.1

    ws.print_options.horizontalCentered = False
    ws.print_options.verticalCentered = False

    if not is_heavy:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1


def process_excel(uploaded_file, lang_code="TH"):
    t = TEXTS[lang_code]
    wb_raw = openpyxl.load_workbook(uploaded_file)
    ws_raw = wb_raw.active
    raw_data = [list(row) for row in ws_raw.iter_rows(values_only=True)]
    raw_df = pd.DataFrame(raw_data)

    header_row_idx = None
    col_mapping = {}

    for r in range(min(15, len(raw_df))):
        row_vals = [
            str(v).strip().lower() if pd.notna(v) else "" for v in raw_df.iloc[r]
        ]

        pid_col = None
        for c, val in enumerate(row_vals):
            if any(k in val for k in ["item pid", "pids", "pid", "code"]):
                pid_col = c
                break

        has_sph = any("sph" in val for val in row_vals)
        has_cyl = any("cyl" in val for val in row_vals)

        if pid_col is not None and (has_sph or has_cyl):
            header_row_idx = r
            col_mapping["pid"] = pid_col

            for c, val in enumerate(row_vals):
                if (
                    any(
                        k in val
                        for k in [
                            "item name",
                            "lk type package",
                            "package",
                            "name",
                            "desc",
                        ]
                    )
                    and c != pid_col
                ):
                    col_mapping["name"] = c
                elif "sph" in val:
                    col_mapping["sph"] = c
                elif "cyl" in val:
                    col_mapping["cyl"] = c
            break

    if "pid" in col_mapping and "name" not in col_mapping:
        col_mapping["name"] = col_mapping["pid"] + 1

    required_cols = ["pid", "name", "sph", "cyl"]
    missing_cols = [c for c in required_cols if c not in col_mapping]
    if header_row_idx is None or len(missing_cols) > 0:
        raise ValueError(
            f"Invalid header format. Missing columns: {', '.join(missing_cols)}"
        )

    item_start_row = header_row_idx + 1
    item_end_row = len(raw_df) - 1
    for r in range(item_start_row, len(raw_df)):
        first_val = str(raw_df.iloc[r, 0]).strip().lower()
        if "grand total" in first_val or "total" in first_val:
            item_end_row = r - 1
            break

    first_store_col = max(col_mapping.values()) + 1
    total_cols = raw_df.shape[1]

    # ตรวจสอบว่ามี Store ID / Store Name อยู่แถวด้านบนจริงหรือไม่
    has_upper_store_info = False
    if header_row_idx >= 2:
        top_row_vals = [
            str(raw_df.iloc[header_row_idx - 2, c]).strip().lower()
            for c in range(first_store_col, total_cols)
            if pd.notna(raw_df.iloc[header_row_idx - 2, c])
        ]
        # ถ้าไม่มีคำว่า 'do' หรือมีรหัสสาขาจริง
        if any(v != "do" and v != "nan" and v != "" for v in top_row_vals):
            has_upper_store_info = True

    default_store_name = uploaded_file.name.replace(".xlsx", "").replace(".XLSX", "")

    store_cols_data = []
    for c_idx in range(first_store_col, total_cols):
        do_val = raw_df.iloc[header_row_idx, c_idx]
        if pd.isna(do_val) or "total" in str(do_val).lower():
            continue

        do_num_str = str(do_val).strip()

        if has_upper_store_info:
            st_id = raw_df.iloc[header_row_idx - 2, c_idx]
            st_name = raw_df.iloc[header_row_idx - 1, c_idx]
            st_id_str = str(st_id).strip() if pd.notna(st_id) else "MAIN"
            st_name_str = str(st_name).strip() if pd.notna(st_name) else st_id_str
        else:
            st_id_str = "00001"
            st_name_str = default_store_name

        store_cols_data.append(
            {
                "col_idx": c_idx,
                "store_id": st_id_str,
                "store_name": st_name_str,
                "do_number": do_num_str,
            }
        )

    stores_df = pd.DataFrame(store_cols_data)

    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)

    thin_black = Side(border_style="thin", color="000000")
    double_black = Side(border_style="double", color="000000")

    box_border = Border(
        left=thin_black, right=thin_black, top=thin_black, bottom=thin_black
    )
    header_border = Border(
        left=thin_black, right=thin_black, top=thin_black, bottom=thin_black
    )
    total_border = Border(
        left=thin_black, right=thin_black, top=thin_black, bottom=double_black
    )

    fill_green = PatternFill(start_color="E2EFDA", fill_type="solid")
    fill_peach = PatternFill(start_color="FCE4D6", fill_type="solid")

    # Tab 1: Summary All Stores
    ws_summary = wb_out.create_sheet(title="Summary All Stores")
    setup_sheet_page_layout(ws_summary, is_heavy=False)

    ws_summary.merge_cells("A1:E1")
    ws_summary["A1"] = "สรุปรายการจัดส่งเลนส์ประจำรอบ (Store Dispatch Summary)"
    ws_summary["A1"].font = Font(
        name="Cordia New", size=18, bold=True, color="000000"
    )
    ws_summary["A1"].alignment = Alignment(
        horizontal="center", vertical="center"
    )

    ws_summary.cell(row=1, column=7, value="🟢 Tab สีเขียว:").font = Font(
        name="Cordia New", size=11, bold=True
    )
    cell_lg_g = ws_summary.cell(
        row=1,
        column=8,
        value="สั่งพิมพ์รวมได้ทันที (ไม่เกิน 8 DO และไม่เกิน 35 แถว)",
    )
    cell_lg_g.font = Font(name="Cordia New", size=11)
    cell_lg_g.fill = fill_green

    ws_summary.cell(row=2, column=7, value="🔴 Tab สีแดง:").font = Font(
        name="Cordia New", size=11, bold=True
    )
    cell_lg_r = ws_summary.cell(
        row=2,
        column=8,
        value="ต้องสั่งพิมพ์แยกทีละ Sheet (เกิน 8 DO หรือเกิน 35 แถว)",
    )
    cell_lg_r.font = Font(name="Cordia New", size=11)
    cell_lg_r.fill = fill_peach

    headers_summary = [
        "ลำดับ",
        "รหัสสาขา (Store ID)",
        "ชื่อสาขา (Store Name)",
        "จำนวน DO",
        "จำนวนเลนส์รวม (ชิ้น)",
    ]
    for col_num, h_text in enumerate(headers_summary, 1):
        cell = ws_summary.cell(row=3, column=col_num, value=h_text)
        cell.font = Font(name="Cordia New", size=13, bold=True, color="000000")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border

    store_groups = list(stores_df.groupby("store_id", sort=False))

    row_idx = 4
    idx = 1
    summary_preview_data = []

    for store_id, group in store_groups:
        store_name = (
            str(group["store_name"].iloc[0])
            if pd.notna(group["store_name"].iloc[0])
            else str(store_id)
        )
        num_dos = len(group)
        st_cols = group["col_idx"].tolist()
        total_pcs = sum(
            int(parse_num(raw_df.iloc[r, c]))
            for r in range(item_start_row, item_end_row + 1)
            for c in st_cols
        )

        active_items_count = sum(
            1
            for r in range(item_start_row, item_end_row + 1)
            if any(parse_num(raw_df.iloc[r, c]) > 0 for c in st_cols)
        )
        is_heavy = active_items_count > 35 or num_dos > 8

        ws_summary.cell(row=row_idx, column=1, value=idx).alignment = Alignment(
            horizontal="center"
        )
        ws_summary.cell(
            row=row_idx, column=2, value=store_id
        ).alignment = Alignment(horizontal="center")
        ws_summary.cell(row=row_idx, column=3, value=store_name)
        ws_summary.cell(
            row=row_idx, column=4, value=num_dos
        ).alignment = Alignment(horizontal="right")
        ws_summary.cell(
            row=row_idx, column=5, value=total_pcs
        ).alignment = Alignment(horizontal="right")

        for c in range(1, 6):
            cell = ws_summary.cell(row=row_idx, column=c)
            cell.font = Font(name="Cordia New", size=12)
            cell.border = box_border

        summary_preview_data.append(
            {
                "Store ID": store_id,
                "Store Name": store_name,
                "DO Count": num_dos,
                "Total Pcs": total_pcs,
                "Status": t["status_heavy"] if is_heavy else t["status_normal"],
            }
        )

        row_idx += 1
        idx += 1

    ws_summary.cell(row=row_idx, column=1, value="Grand Total").font = Font(
        name="Cordia New", size=12, bold=True
    )
    ws_summary.cell(row=row_idx, column=1).alignment = Alignment(
        horizontal="center"
    )
    ws_summary.cell(
        row=row_idx, column=4, value=f"=SUM(D4:D{row_idx-1})"
    ).font = Font(name="Cordia New", size=12, bold=True)
    ws_summary.cell(row=row_idx, column=4).alignment = Alignment(
        horizontal="right"
    )
    ws_summary.cell(
        row=row_idx, column=5, value=f"=SUM(E4:E{row_idx-1})"
    ).font = Font(name="Cordia New", size=12, bold=True)
    ws_summary.cell(row=row_idx, column=5).alignment = Alignment(
        horizontal="right"
    )

    for c in range(1, 6):
        cell = ws_summary.cell(row=row_idx, column=c)
        cell.border = total_border

    ws_summary.column_dimensions["A"].width = 8
    ws_summary.column_dimensions["B"].width = 22
    ws_summary.column_dimensions["C"].width = 38
    ws_summary.column_dimensions["D"].width = 16
    ws_summary.column_dimensions["E"].width = 22

    priority_stores = []
    normal_stores = []

    for store_id, group in store_groups:
        st_cols = group["col_idx"].tolist()
        num_dos = len(st_cols)
        active_items_count = sum(
            1
            for r in range(item_start_row, item_end_row + 1)
            if any(parse_num(raw_df.iloc[r, c]) > 0 for c in st_cols)
        )

        if active_items_count > 35 or num_dos > 8:
            priority_stores.append((store_id, group, True, num_dos))
        else:
            normal_stores.append((store_id, group, False, num_dos))

    priority_stores_sorted = sorted(priority_stores, key=lambda x: x[3])
    normal_stores_sorted = sorted(normal_stores, key=lambda x: x[3])
    sorted_store_groups = priority_stores_sorted + normal_stores_sorted

    single_store_files = {}

    for store_id, group, is_heavy, _ in sorted_store_groups:
        store_name = (
            str(group["store_name"].iloc[0])
            if pd.notna(group["store_name"].iloc[0])
            else str(store_id)
        )
        sheet_title = f"{store_id} - {store_name}"[:28]

        store_cols_all = group["col_idx"].tolist()
        do_nums_all = group["do_number"].tolist()

        dos_with_counts = []
        for c_idx, do_n in zip(store_cols_all, do_nums_all):
            cnt = sum(
                1
                for r in range(item_start_row, item_end_row + 1)
                if parse_num(raw_df.iloc[r, c_idx]) > 0
            )
            dos_with_counts.append((c_idx, do_n, cnt))

        dos_sorted_by_count = sorted(dos_with_counts, key=lambda x: x[2])
        ordered_dos = [(x[0], x[1]) for x in dos_sorted_by_count]

        ordered_do_nums = [x[1] for x in ordered_dos]
        num_dos = len(ordered_do_nums)

        chunk_size = 8
        do_chunks = [
            ordered_dos[i : i + chunk_size]
            for i in range(0, num_dos, chunk_size)
        ]

        ws = wb_out.create_sheet(title=sheet_title)
        setup_sheet_page_layout(ws, is_heavy=is_heavy)

        if is_heavy:
            ws.sheet_properties.tabColor = "FCE4D6"
        else:
            ws.sheet_properties.tabColor = "E2EFDA"

        curr_row = 1
        total_chunks = len(do_chunks)
        global_do_counter = 1

        for chunk_idx, chunk_dos in enumerate(do_chunks):
            chunk_do_nums = [x[1] for x in chunk_dos]
            chunk_num_dos = len(chunk_dos)
            is_last_chunk = chunk_idx == (total_chunks - 1)

            chunk_items_list = []
            for sub_do_idx, (c_idx, do_n) in enumerate(chunk_dos):
                do_items = []
                for r_idx in range(item_start_row, item_end_row + 1):
                    q_val = parse_num(raw_df.iloc[r_idx, c_idx])
                    if q_val > 0:
                        pid_val = raw_df.iloc[r_idx, col_mapping["pid"]]
                        name_val = raw_df.iloc[r_idx, col_mapping["name"]]
                        sph_val = parse_num(
                            raw_df.iloc[r_idx, col_mapping["sph"]]
                        )
                        cyl_val = parse_num(
                            raw_df.iloc[r_idx, col_mapping["cyl"]]
                        )

                        do_items.append(
                            {
                                "pid": (
                                    str(pid_val) if pd.notna(pid_val) else ""
                                ),
                                "name": (
                                    str(name_val) if pd.notna(name_val) else ""
                                ),
                                "sph": sph_val,
                                "cyl": cyl_val,
                                "qty": (
                                    int(q_val) if q_val == int(q_val) else q_val
                                ),
                                "sub_do_idx": sub_do_idx,
                            }
                        )

                do_items_sorted = sorted(
                    do_items,
                    key=lambda x: (x["name"], -x["sph"], -x["cyl"], x["pid"]),
                )

                for item in do_items_sorted:
                    qty_array = [None] * chunk_num_dos
                    qty_array[item["sub_do_idx"]] = item["qty"]
                    chunk_items_list.append(
                        {
                            "pid": item["pid"],
                            "name": item["name"],
                            "sph": item["sph"],
                            "cyl": item["cyl"],
                            "qtys": qty_array,
                        }
                    )

            if chunk_idx > 0:
                ws.row_breaks.append(Break(id=curr_row - 1))

            ws.cell(
                row=curr_row, column=1, value=f"Store ID: {store_id}"
            ).font = Font(name="Cordia New", size=11, bold=True)
            ws.cell(
                row=curr_row, column=3, value=f"Store Name: {store_name}"
            ).font = Font(name="Cordia New", size=11, bold=True)
            curr_row += 1

            for idx_q in range(chunk_num_dos):
                c_i = 5 + idx_q
                cell_seq = ws.cell(
                    row=curr_row, column=c_i, value=global_do_counter + idx_q
                )
                cell_seq.font = Font(
                    name="Cordia New", size=11, bold=True, color="000000"
                )
                cell_seq.alignment = Alignment(
                    horizontal="center", vertical="center"
                )
                cell_seq.border = header_border

            curr_row += 1

            base_headers = ["Item PID", "Item Name", "SPH", "CYL"]
            for col_i, h_text in enumerate(base_headers, 1):
                cell = ws.cell(row=curr_row, column=col_i, value=h_text)
                cell.font = Font(
                    name="Cordia New", size=11, bold=True, color="000000"
                )
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )
                cell.border = header_border

            for idx_q, do_n in enumerate(chunk_do_nums):
                c_i = 5 + idx_q
                cell = ws.cell(row=curr_row, column=c_i, value=f"DO: {do_n}")
                cell.font = Font(
                    name="Cordia New", size=11, bold=True, color="000000"
                )
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )
                cell.border = header_border

            global_do_counter += chunk_num_dos
            max_col_idx = 4 + chunk_num_dos
            curr_row += 1
            start_data_row = curr_row

            for row_data in chunk_items_list:
                row_data_row = curr_row
                ws.cell(
                    row=row_data_row, column=1, value=row_data["pid"]
                ).alignment = Alignment(horizontal="center")
                ws.cell(
                    row=row_data_row, column=2, value=row_data["name"]
                ).alignment = Alignment(horizontal="left")

                sph_fmt = (
                    f"{row_data['sph']:+.2f}" if row_data["sph"] != 0 else "0.00"
                )
                cyl_fmt = (
                    f"{row_data['cyl']:+.2f}" if row_data["cyl"] != 0 else "0.00"
                )

                ws.cell(
                    row=row_data_row, column=3, value=sph_fmt
                ).alignment = Alignment(horizontal="right")
                ws.cell(
                    row=row_data_row, column=4, value=cyl_fmt
                ).alignment = Alignment(horizontal="right")

                for idx_q, q_val in enumerate(row_data["qtys"]):
                    c_i = 5 + idx_q
                    ws.cell(
                        row=row_data_row, column=c_i, value=q_val
                    ).alignment = Alignment(horizontal="right")

                for c in range(1, max_col_idx + 1):
                    cell = ws.cell(row=row_data_row, column=c)
                    cell.font = Font(name="Cordia New", size=11)
                    cell.border = box_border

                curr_row += 1

            end_data_row = curr_row - 1

            if is_last_chunk:
                ws.cell(
                    row=curr_row, column=1, value="Grand Total"
                ).font = Font(name="Cordia New", size=11, bold=True)
                ws.cell(row=curr_row, column=1).alignment = Alignment(
                    horizontal="center"
                )

                for idx_q in range(chunk_num_dos):
                    col_idx = 5 + idx_q
                    col_letter = get_column_letter(col_idx)

                    if start_data_row <= end_data_row:
                        sum_formula = f"=SUM({col_letter}{start_data_row}:{col_letter}{end_data_row})"
                    else:
                        sum_formula = 0

                    ws.cell(
                        row=curr_row, column=col_idx, value=sum_formula
                    ).font = Font(name="Cordia New", size=11, bold=True)
                    ws.cell(row=curr_row, column=col_idx).alignment = Alignment(
                        horizontal="right"
                    )

                for c in range(1, max_col_idx + 1):
                    cell = ws.cell(row=curr_row, column=c)
                    cell.border = total_border

                curr_row += 1

            curr_row += 1

        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 9
        ws.column_dimensions["D"].width = 9
        for idx_q in range(min(8, num_dos)):
            ws.column_dimensions[get_column_letter(5 + idx_q)].width = 13

        wb_single = openpyxl.Workbook()
        wb_single.remove(wb_single.active)
        ws_single = wb_single.create_sheet(title=sheet_title)

        for row in ws.iter_rows(values_only=False):
            for cell in row:
                ws_single.cell(
                    row=cell.row, column=cell.column, value=cell.value
                )

        setup_sheet_page_layout(ws_single, is_heavy=is_heavy)
        single_buf = io.BytesIO()
        wb_single.save(single_buf)
        single_buf.seek(0)
        single_store_files[f"{store_id}_{store_name}"] = single_buf

    for sheet in wb_out.worksheets:
        sheet.sheet_view.tabSelected = True

    output = io.BytesIO()
    wb_out.save(output)
    output.seek(0)

    stats = {
        "total_stores": len(store_groups),
        "heavy_stores": len(priority_stores),
        "normal_stores": len(normal_stores),
        "total_pcs": sum(item["Total Pcs"] for item in summary_preview_data),
        "preview_df": pd.DataFrame(summary_preview_data),
        "single_files": single_store_files,
    }

    return output, stats


# ==============================================================================
# 4. ส่วนจัดวางหน้าตาเว็บ (Multi-language Layout)
# ==============================================================================

col_space, col_lang = st.columns([5, 1])
with col_lang:
    selected_lang = st.selectbox(
        "🌐 Language / ภาษา",
        ["🇹🇭 ไทย", "🇬🇧 English"],
        index=0,
        label_visibility="collapsed",
    )

lang_code = "TH" if "ไทย" in selected_lang else "EN"
t = TEXTS[lang_code]

if "processed_results" not in st.session_state:
    st.session_state["processed_results"] = None

if st.session_state["processed_results"] is None:
    _, center_col, _ = st.columns([1, 2.2, 1])

    with center_col:
        st.markdown(
            f"""
            <div class="header-box">
                <div class="header-title">{t['title']}</div>
                <div class="header-subtitle">{t['subtitle']}</div>
            </div>
        """,
            unsafe_allow_html=True,
        )

        uploaded_files = st.file_uploader(
            t["uploader_label"],
            type=["xlsx"],
            accept_multiple_files=True,
        )

        with st.expander(t["guide_title"]):
            st.markdown(t["guide_content"], unsafe_allow_html=True)

        if uploaded_files:
            if st.button(t["btn_process"]):
                processed_results = []
                errors = []

                progress_bar = st.progress(0)
                status_text = st.empty()

                for i, file in enumerate(uploaded_files):
                    status_text.text(
                        f"⏳ Processing ({i+1}/{len(uploaded_files)}): {file.name}"
                    )
                    try:
                        out_bytes, stats = process_excel(
                            file, lang_code=lang_code
                        )
                        processed_results.append((file.name, out_bytes, stats))
                    except Exception as e:
                        errors.append((file.name, str(e)))

                    progress_bar.progress((i + 1) / len(uploaded_files))

                status_text.empty()

                if errors:
                    for err_file, err_msg in errors:
                        st.error(f"❌ `{err_file}`: {err_msg}")

                if processed_results:
                    st.session_state["processed_results"] = processed_results
                    st.rerun()

else:
    processed_results = st.session_state["processed_results"]

    st.markdown(
        f"""
        <div class="header-box">
            <div class="header-title">{t['title']}</div>
            <div class="header-subtitle">{t['subtitle']}</div>
        </div>
    """,
        unsafe_allow_html=True,
    )

    st.success(t["success"])

    _, _, first_stats = processed_results[0]
    df_preview = first_stats["preview_df"]

    st.markdown(f"### {t['dashboard_title']}")
    m1, m2, m3, m4 = st.columns(4)
    m1.metric(
        t["m_total_stores"],
        f"{first_stats['total_stores']} {t['unit_store']}",
    )
    m2.metric(
        t["m_heavy_stores"],
        f"{first_stats['heavy_stores']} {t['unit_store']}",
    )
    m3.metric(
        t["m_normal_stores"],
        f"{first_stats['normal_stores']} {t['unit_store']}",
    )
    m4.metric(
        t["m_total_pcs"], f"{first_stats['total_pcs']:,} {t['unit_pcs']}"
    )

    st.markdown("<br>", unsafe_allow_html=True)

    col_search, col_filter = st.columns([2, 1])
    search_query = col_search.text_input(t["search_label"], "")
    filter_status = col_filter.selectbox(
        t["filter_label"],
        [t["filter_all"], t["filter_normal"], t["filter_heavy"]],
    )

    filtered_df = df_preview.copy()
    if search_query:
        filtered_df = filtered_df[
            filtered_df["Store ID"]
            .str.contains(search_query, case=False, na=False)
            | filtered_df["Store Name"].str.contains(
                search_query, case=False, na=False
            )
        ]

    if filter_status == t["filter_normal"]:
        filtered_df = filtered_df[
            filtered_df["Status"].str.contains("Normal")
        ]
    elif filter_status == t["filter_heavy"]:
        filtered_df = filtered_df[filtered_df["Status"].str.contains("Heavy")]

    st.dataframe(filtered_df, use_container_width=True)

    st.markdown("<hr>", unsafe_allow_html=True)

    col_download, col_single = st.columns([1, 1])

    with col_download:
        st.markdown(f"#### {t['main_export']}")
        if len(processed_results) == 1:
            fname, fbytes, _ = processed_results[0]
            out_name = f"Consolidated_{fname}"
            st.download_button(
                label=f"{t['btn_download_excel']} ({out_name})",
                data=fbytes,
                file_name=out_name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        else:
            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(
                zip_buffer, "w", zipfile.ZIP_DEFLATED
            ) as zip_file:
                for fname, fbytes, _ in processed_results:
                    zip_file.writestr(
                        f"Consolidated_{fname}", fbytes.getvalue()
                    )

            zip_buffer.seek(0)
            st.download_button(
                label=t["btn_download_zip"],
                data=zip_buffer,
                file_name="All_Consolidated_Lists.zip",
                mime="application/zip",
            )

    with col_single:
        st.markdown(f"#### {t['single_export']}")
        selected_store_key = st.selectbox(
            t["select_single"],
            options=list(first_stats["single_files"].keys()),
        )
        if selected_store_key:
            st.download_button(
                label=f"{t['btn_single']} ({selected_store_key})",
                data=first_stats["single_files"][selected_store_key],
                file_name=f"Dispatch_{selected_store_key}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

    st.markdown("<br>", unsafe_allow_html=True)
    if st.button(t["btn_reset"]):
        st.session_state["processed_results"] = None
        st.rerun()
